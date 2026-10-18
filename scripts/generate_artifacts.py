import json
from pathlib import Path
from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from app.common import PREFIX
from app.vulnerabilities import app as vapp
from app.notifications import app as napp
from app.playbooks import app as papp
from app.common import app as authapp

V={'title':'Exposed remote service','description':'Scanner detected an exposed service requiring review.','cve_identifier':'CVE-2026-12345','cvss_score':9.1,'severity':'Critical','asset_id':1,'discovery_source':'Lab scanner'}
N={'user_id':3,'title':'Security alert','message':'Review the reported suspicious message.','notification_type':'Security Alert','severity':'High'}
P={'name':'Malware containment','description':'Isolate and investigate a suspected malware infection.','incident_category':'Malware','severity':'High'}

items=[]
folder=None

def group(name):
    global folder
    folder={'name':name,'item':[]}
    items.append(folder)

def add(name,port,method,path,body=None,role='admin',expect=200,save=None,description=''):
    url='{{base'+str(port)+'}}'+PREFIX+path
    headers=[{'key':'Content-Type','value':'application/json'}]
    if role:
        headers.append({'key':'Authorization','value':'Bearer {{'+role+'_token}}'})
    req={'method':method,'header':headers,'url':url,'description':description or name+'. Actual response is checked by the collection test.'}
    if body is not None:
        req['body']={'mode':'raw','raw':json.dumps(body,indent=2),'options':{'raw':{'language':'json'}}}
    tests=[f'pm.test("HTTP {expect}", function () {{ pm.response.to.have.status({expect}); }});']
    if expect >= 400:
        tests.append('pm.test("Consistent error envelope", function () { const body = pm.response.json(); pm.expect(body.status).to.eql("error"); pm.expect(body.message).to.be.a("string"); });')
    elif expect == 201:
        tests.append('pm.test("Persisted record identifier", function () { pm.expect(pm.response.json().id).to.be.a("number"); });')
    if body and body.get('status') in ('Resolved','Completed'):
        key = 'resolution_date' if '/vulnerabilities/' in path else 'completion_date'
        tests.append(f'pm.test("Server recorded completion timestamp", function () {{ pm.expect(pm.response.json().{key}).to.be.a("string"); }});')

    if save:
        variable,key=save
        tests.append(f'if (pm.response.code === {expect}) pm.environment.set("{variable}", pm.response.json().{key});')
    folder['item'].append({'name':name,'request':req,'event':[{'listen':'test','script':{'type':'text/javascript','exec':tests}}]})

def generate():
    items.clear()
    group('01 Authentication')
    for role in ('admin','analyst','employee','auditor'):
        add('Login '+role,8000,'POST','/auth/login',{'username':role,'password':'{{'+role+'_password}}'},None,save=(role+'_token','access_token'))
    add('Reject incorrect password',8000,'POST','/auth/login',{'username':'admin','password':'incorrect'},None,401)
    group('02 Vulnerability management')
    add('Create critical vulnerability',8007,'POST','/vulnerabilities',V,expect=201,save=('vulnerability_id','id'))
    add('Create second vulnerability',8007,'POST','/vulnerabilities',V|{'title':'Workstation review','asset_id':2,'severity':'Medium','cvss_score':5.5,'cve_identifier':'CVE-2026-23456'},expect=201)
    add('Retrieve vulnerabilities',8007,'GET','/vulnerabilities')
    add('Retrieve vulnerability',8007,'GET','/vulnerabilities/{{vulnerability_id}}')
    add('Missing vulnerability',8007,'GET','/vulnerabilities/99999999',expect=404)
    for query in ('severity=Critical','asset_id=2','cve_identifier=CVE-2026-12345','status=Open'):
        add('Search '+query,8007,'GET','/vulnerabilities/search?'+query,role='analyst')
    add('Reject CVSS above ten',8007,'POST','/vulnerabilities',V|{'cvss_score':11},expect=422)
    add('Reject missing title',8007,'POST','/vulnerabilities',{k:v for k,v in V.items() if k!='title'},expect=422)
    add('Reject employee access',8007,'GET','/vulnerabilities',role='employee',expect=403)
    add('Reject missing token',8007,'GET','/vulnerabilities',role=None,expect=401)
    add('Reject lifecycle skip',8007,'PUT','/vulnerabilities/{{vulnerability_id}}',{'status':'Closed'},expect=409)
    for state in ('Under Investigation','In Progress','Resolved','Closed'):
        add('Advance vulnerability to '+state,8007,'PUT','/vulnerabilities/{{vulnerability_id}}',{'status':state},role='analyst')
    add('Reject analyst deletion',8007,'DELETE','/vulnerabilities/{{vulnerability_id}}',role='analyst',expect=403)
    add('Delete vulnerability',8007,'DELETE','/vulnerabilities/{{vulnerability_id}}')
    group('03 Notifications')
    add('Create employee alert',8008,'POST','/notifications',N,expect=201,save=('notification_id','id'))
    add('Create analyst alert',8008,'POST','/notifications',N|{'user_id':2,'notification_type':'Incident Assignment','title':'Incident assigned'},expect=201)
    add('Retrieve own notifications',8008,'GET','/notifications/user/3',role='employee')
    add('Retrieve empty notifications',8008,'GET','/notifications/user/4')
    add('Retrieve specific notification',8008,'GET','/notifications/{{notification_id}}',role='employee')
    add('Reject another user notification',8008,'GET','/notifications/{{notification_id}}',role='analyst',expect=403)
    add('Missing notification',8008,'GET','/notifications/99999999',expect=404)
    for query in ('user_id=3','notification_type=Security%20Alert','severity=High','status=Unread'):
        add('Search '+query,8008,'GET','/notifications/search?'+query,role='auditor')
    add('Reject auditor marking read',8008,'PUT','/notifications/{{notification_id}}/read',role='auditor',expect=403)
    add('Mark own notification read',8008,'PUT','/notifications/{{notification_id}}/read',role='employee')
    add('Repeat mark read',8008,'PUT','/notifications/{{notification_id}}/read',role='employee')
    add('Reject employee deletion',8008,'DELETE','/notifications/{{notification_id}}',role='employee',expect=403)
    add('Delete notification',8008,'DELETE','/notifications/{{notification_id}}')
    group('04 Playbooks and ordered actions')
    add('Create playbook',8009,'POST','/playbooks',P,expect=201,save=('playbook_id','id'))
    add('Create second playbook',8009,'POST','/playbooks',P|{'name':'Phishing investigation','incident_category':'Phishing','severity':'Medium'},expect=201)
    add('List playbooks',8009,'GET','/playbooks',role='auditor')
    add('Retrieve playbook',8009,'GET','/playbooks/{{playbook_id}}',role='analyst')
    add('Missing playbook',8009,'GET','/playbooks/99999999',expect=404)
    for step in (1,2):
        add('Add response step '+str(step),8009,'POST','/playbooks/{{playbook_id}}/actions',{'description':['Isolate asset','Investigate artifacts'][step-1],'step_number':step,'assigned_user_id':2},role='analyst',expect=201,save=('action'+str(step)+'_id','id'))
    add('Reject duplicate step',8009,'POST','/playbooks/{{playbook_id}}/actions',{'description':'Duplicate','step_number':1},expect=409)
    add('Reject out of order action',8009,'PUT','/actions/{{action2_id}}/status',{'status':'In Progress'},role='analyst',expect=409)
    for step in (1,2):
        for state in ('In Progress','Completed'):
            add('Step '+str(step)+' '+state,8009,'PUT','/actions/{{action'+str(step)+'_id}}/status',{'status':state},role='analyst')
    add('Retrieve completed workflow',8009,'GET','/playbooks/{{playbook_id}}')
    add('Reject analyst deleting playbook',8009,'DELETE','/playbooks/{{playbook_id}}',role='analyst',expect=403)
    add('Delete playbook',8009,'DELETE','/playbooks/{{playbook_id}}')
    collection={'info':{'name':'Joseph Awaida — Services 7, 8, 9','schema':'https://schema.getpostman.com/json/collection/v2.1.0/collection.json','description':'Run folders in order. Set the four local demo passwords in the environment. Tokens and created IDs populate automatically. Includes successful, missing, invalid and forbidden cases.'},'item':items}
    Path('postman/Services-7-8-9.postman_collection.json').write_text(json.dumps(collection,indent=2))
    values=[{'key':'base'+str(port),'value':'http://127.0.0.1:'+str(port),'enabled':True} for port in (8000,8007,8008,8009)]
    values += [{'key':role+'_password','value':'','type':'secret','enabled':True} for role in ('admin','analyst','employee','auditor')]
    values += [{'key':role+'_token','value':'','type':'secret','enabled':True} for role in ('admin','analyst','employee','auditor')]
    Path('postman/Local.postman_environment.json').write_text(json.dumps({'name':'Joseph Awaida Local','values':values,'_postman_variable_scope':'environment'},indent=2))
    text=['# API reference (Joseph Awaida)','All protected requests require a bearer token. Full request/response schemas are in the exported OpenAPI JSON and interactive `/docs`. Error format: `{"status":"error","message":"..."}`. Common errors: 401 missing/invalid authentication; 403 role/ownership; 404 missing resource; 409 conflicting step or workflow; 422 input validation.','']
    for app in (authapp,vapp,napp,papp):
        spec=app.openapi()
        Path('docs/'+app.title.split(':')[0].replace(' ','-').replace('(','').replace(')','')+'.openapi.json').write_text(json.dumps(spec,indent=2))
        text += ['## '+app.title,'']
        for path,methods in spec['paths'].items():
            for method,op in methods.items():
                text += ['### '+op['summary'],f'`{method.upper()} {path}`','']
                if 'parameters' in op:
                    text += ['Parameters: '+', '.join(p['name']+' ('+p['in']+', '+('required' if p.get('required') else 'optional')+')' for p in op['parameters']), '']
                if 'requestBody' in op:
                    ref=op['requestBody']['content']['application/json']['schema'].get('$ref','')
                    schema=spec['components']['schemas'].get(ref.split('/')[-1],{})
                    text += ['JSON body: `'+', '.join(schema.get('properties',{}))+'`. Required: `'+', '.join(schema.get('required',[]))+'`.','']
                if op.get('description'):
                    text += [op['description'], '']
                text += ['Success: '+', '.join(code for code in op['responses'] if code.startswith('2'))+'. See Postman requests for runnable examples.','']
                for code, response in op['responses'].items():
                    if not code.startswith('2'):
                        continue
                    schema = response.get('content', {}).get('application/json', {}).get('schema', {})
                    array = schema.get('type') == 'array'
                    ref = schema.get('items', {}).get('$ref', '') if array else schema.get('$ref', '')
                    name = ref.rsplit('/',1)[-1]
                    fields = spec.get('components', {}).get('schemas', {}).get(name, {}).get('properties', {})
                    if fields:
                        text += ['Response: ' + ('array of ' if array else '') + name + ' objects with fields `' + ', '.join(fields) + '`.','']

    Path('docs/API.md').write_text('\n'.join(text))
    doc=Document()
    doc.add_heading('Security Incident Detection and Response Management System',0)
    doc.add_paragraph('Sprint I — Services 7, 8, 9\nJoseph Awaida\n18 October 2026\nContribution report: merge with the common team report before submission.')
    doc.add_heading('Table of contents',1)
    p=doc.add_paragraph()
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'TOC \\o "1-3" \\h \\z \\u'); p._p.append(field)
    doc.add_paragraph('Update this table in Word to populate page numbers. All sections in this contribution: Joseph Awaida.')
    for line in Path('docs/REPORT.md').read_text().splitlines():
        if line.startswith('# '):
            continue
        if line.startswith('## '):
            doc.add_heading(line[3:],1)
        elif line.startswith('### '):
            doc.add_heading(line[4:],2)
        elif line and not line.startswith('!['):
            doc.add_paragraph(line)
    evidence_file = Path('evidence/http-execution.json')
    if evidence_file.exists():
        records = json.loads(evidence_file.read_text())
        doc.add_heading('Actual execution records (Joseph Awaida)',1)
        doc.add_paragraph('Real HTTP requests captured at ' + records['captured_at'] + '. These records supplement required Postman screenshots; they are not screenshots.')
        table = doc.add_table(rows=1, cols=3)
        table.style = 'Light Shading Accent 1'
        for cell, value in zip(table.rows[0].cells, ['Operation','Expected','Actual']):
            cell.text = value
        for record in records['results']:
            for cell, value in zip(table.add_row().cells, [record['name'], str(record['expected_status']), str(record['actual_status'])]):
                cell.text = value
        for name in ['Create critical vulnerability','Create employee alert','Retrieve completed workflow']:
            record = next(r for r in records['results'] if r['name'] == name)
            doc.add_heading(name,2)
            doc.add_paragraph(record['method'] + ' ' + record['url'])
            doc.add_paragraph(json.dumps(record['response'],indent=2))
        if Path('evidence/newman-output.txt').exists():
            doc.add_heading('Postman Newman validation',2)
            summary_path = Path('evidence/validation-summary.json')
            if summary_path.exists():
                summary = json.loads(summary_path.read_text())
                n = summary['newman']; p = summary['pytest']
                doc.add_paragraph(f"Actual Newman validation: {n['requests']['total']} requests, {n['assertions']['total']} assertions, {n['assertions']['failed']} failed assertions. Pytest: {p['tests']} tests, {p['failures']} failures, {p['errors']} errors. Full output: evidence/.")

    transitions_path = Path('evidence/database-transitions.json')
    if transitions_path.exists():
        doc.add_heading('Actual before/after database records (Joseph Awaida)',1)
        doc.add_paragraph('Queried immediately before and after the named operation. These records supplement required database screenshots.')
        records = json.loads(transitions_path.read_text())
        for name in ['Advance vulnerability to Resolved','Mark own notification read','Step 1 Completed']:
            record = next(r for r in records if r['operation'] == name)
            doc.add_heading(name,2)
            columns = ['id','status'] + (['read_date'] if record['table']=='notifications' else ['completion_date'] if record['table']=='actions' else ['resolution_date'])
            for phase in ['before','after']:
                doc.add_paragraph(phase.capitalize() + ' — ' + record['table'])
                table = doc.add_table(rows=1, cols=len(columns))
                table.style = 'Light Shading Accent 1'
                for cell,value in zip(table.rows[0].cells, columns):
                    cell.text = value
                for row in record[phase]:
                    for cell,col in zip(table.add_row().cells, columns):
                        cell.text = str(row.get(col))
    integration_path = Path('evidence/http-integration.json')
    if integration_path.exists():
        integration = json.loads(integration_path.read_text())
        doc.add_heading('Actual Service 7 to Service 8 HTTP verification (Joseph Awaida)',1)
        doc.add_paragraph('A vulnerability creation queued a real owner alert; HTTP delivery succeeded. Repeating the same event returned the same notification identifier: ' + str(integration['duplicate_notification_id']) + '. No duplicate alert was created.')
    doc.add_heading('API reference (Joseph Awaida)',1)
    for line in Path('docs/API.md').read_text().splitlines():
        if line.startswith('### '):
            doc.add_heading(line[4:],3)
        elif line.startswith('## '):
            doc.add_heading(line[3:],2)
        elif line and not line.startswith('# '):
            doc.add_paragraph(line)
    doc.add_heading('Design diagrams (Joseph Awaida)',1)
    for name, caption in [('architecture','Service architecture: local processes and team integration boundaries.'),('schema','Database design: reference entities and owned service tables.')]:
        image = Path('docs') / (name + '.png')
        if image.exists():
            doc.add_picture(str(image), width=Inches(6))
            doc.add_paragraph(caption)
    doc.save('docs/Joseph-Awaida-Sprint-I-Report.docx')
    print(f'Generated {sum(len(f["item"]) for f in items)} Postman requests, OpenAPI exports, API reference, and DOCX report.')

if __name__=='__main__':
    generate()
