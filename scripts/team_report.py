"""Build the common team report from merged source and genuine execution evidence."""
import importlib
import json
import re
import subprocess
from html import escape
from pathlib import Path
from PIL import Image as PILImage, ImageDraw
from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.platypus import BaseDocTemplate,PageTemplate,Frame,Paragraph,Spacer,PageBreak,Image,LongTable,KeepTogether
from reportlab.platypus.tableofcontents import TableOfContents
from scripts.diagrams import font

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs/team'
NAMES='Hassan Nasrallah, Mohammad Nasrallah, Joseph Awaida'
TEAM_NAME='Hassan-Nasrallah-Mohammad-Nasrallah-Joseph-Awaida-430L'
SERVICES=[(1,'Users and Identity Management','Hassan Nasrallah','users',8000),
          (2,'Asset Management','Hassan Nasrallah','assets',8002),(3,'Incident Management','Hassan Nasrallah','incidents',8003),
          (4,'Threat Intelligence','Mohammad Nasrallah','threats',8004),(5,'Security Audit Logging','Mohammad Nasrallah','logs',8005),
          (6,'Security Reports','Mohammad Nasrallah','reports',8006),(7,'Vulnerability Management','Joseph Awaida','vulnerabilities',8007),
          (8,'Notification and Alert Management','Joseph Awaida','notifications',8008),(9,'Security Playbooks and Response Automation','Joseph Awaida','playbooks',8009)]


def diagrams():
    image=PILImage.new('RGB',(1600,1120),'#f7fafc');d=ImageDraw.Draw(image)
    d.text((30,25),'Integrated nine-service architecture',font=font(32),fill='#183b56')
    d.text((30,70),NAMES,font=font(20),fill='#526879')
    d.rounded_rectangle((410,120,1190,200),radius=12,fill='#dcebf6',outline='#38647e',width=2)
    d.text((490,140),'Postman / REST clients — shared bearer sessions',font=font(24),fill='#183b56')
    for i,(number,title,owner,module,port) in enumerate(SERVICES):
        col=i%3;row=i//3;x=35+col*525;y=280+row*220
        d.line((800,200,x+245,y),fill='#a0b4c4',width=2)
        d.rounded_rectangle((x,y,x+500,y+160),radius=12,fill='white',outline='#38647e',width=2)
        d.text((x+15,y+15),f'Service {number} — :{port}',font=font(25),fill='#183b56')
        display=title if len(title)<34 else 'Playbooks and response automation'
        d.text((x+15,y+60),display,font=font(21),fill='#294f69')
        d.text((x+15,y+108),owner,font=font(20),fill='#526879')
    d.rounded_rectangle((150,970,1450,1080),radius=12,fill='#dcebf6',outline='#38647e',width=2)
    d.text((190,988),'Shared SQLite: identities, assets, incidents, threats, logs, vulnerabilities, notifications, playbooks',font=font(20),fill='#183b56')
    d.text((190,1023),'HTTP: report collection from source APIs; queued incident/vulnerability alerts to Service 8',font=font(20),fill='#183b56')
    image.save(OUT/'architecture.png')
    relationships=[('Users','Assets','one owner to many assets'),('Assets','Incidents','one asset to many incidents'),('Users','Incidents','reporter and assigned analyst'),('Incidents','Investigation notes','one incident to many notes'),('Assets','Vulnerabilities','one asset to many vulnerabilities'),('Users','Notifications','one recipient to many alerts'),('Users','Audit logs','one actor to many records'),('Users','Playbooks','creator relationship'),('Playbooks','Response actions','one playbook to ordered steps'),('Users','Response actions','assigned analyst/administrator')]
    image=PILImage.new('RGB',(1600,1150),'#f7fafc');d=ImageDraw.Draw(image)
    d.text((30,25),'Logical database relationships',font=font(32),fill='#183b56')
    d.text((30,75),'Actual columns, primary keys and foreign keys are listed in the schema appendix.',font=font(20),fill='#526879')
    for i,(left,right,label) in enumerate(relationships):
        y=140+i*94
        for x,text in [(40,left),(1130,right)]:
            d.rounded_rectangle((x,y,x+420,y+65),radius=8,fill='white',outline='#38647e',width=2)
            d.text((x+18,y+18),text,font=font(22),fill='#183b56')
        d.line((465,y+33,1120,y+33),fill='#38647e',width=3)
        d.text((485,y+8),label,font=font(19),fill='#294f69')
    d.text((40,1105),'Extension tables and retained histories avoid destructive migration or loss of audit evidence.',font=font(19),fill='#526879')
    image.save(OUT/'database-relationships.png')


def narrative():
    validation=json.loads((ROOT/'evidence/team/validation-summary.json').read_text())
    h=(ROOT/'docs/HASSAN-SPRINT1.md').read_text()
    m=(ROOT/'docs/MOHAMMAD-SPRINT1.md').read_text()
    j=(ROOT/'docs/REPORT.md').read_text()
    h=h[h.index('## Authentication and permissions'):h.index('## Verification and required evidence')]
    m=m[m.index('## Authentication and permissions'):m.index('## Verification and required evidence')]
    j=j[j.index('## 5. Implementation details'):j.index('## 6. Database design')]
    text=f'''# Security Incident Detection and Response Management System — Sprint I

## Title page ({NAMES})
Software Tools Lab, Fall 2026–2027. Team: Hassan Nasrallah (Services 1–3), Mohammad Nasrallah (Services 4–6), Joseph Awaida (Services 7–9). Common team report assembled on 20 October 2026. Repository: https://github.com/j123awaida/435L. Submission folder: {TEAM_NAME}.

## Introduction ({NAMES})
The backend records cybersecurity incidents, organizational assets, threat indicators, vulnerabilities, response workflows and audit activity. It provides nine independent REST services with authenticated access, four role types and persistent relational storage. A frontend is outside Sprint I scope. The goals are accountable incident investigation, repeatable response steps, timely user alerts and security reports derived from current records.

## System architecture ({NAMES})
Each service runs in its own FastAPI process. Hassan's Service 1 issues shared opaque bearer tokens; all protected services verify the session hash, expiry, current account status and role. The integrated launcher is scripts.run_team. The default service ports are 8000 and 8002–8009. The final validation used an isolated database on ports 18000 and 18002–18009 to preserve existing local records.
All services share SQLite. Hassan's schema adds user/asset profile tables and incident/note/history tables to Joseph's compatible reference identities/assets. Mohammad adds threat records and audit-ingestion/retention tables. The application initializes all three schema files in order. This is a local Sprint I architecture with separate API processes, rather than independently deployed database servers.
Incident assignments and vulnerability discovery transactionally queue notifications; a worker authenticates through Hassan's real login endpoint and delivers alerts to Joseph's Service 8 over HTTP. Durable retries and idempotency keys prevent duplicate notification delivery. Mohammad's Service 6 collects user, asset, incident, threat, audit and vulnerability records through the owning services' HTTP APIs. Source metadata reports whether HTTP or database fallback was used; final verification required every source to use HTTP and checked the returned figures against the database.
Audit records already written by the shared middleware and authorization checks are immediately visible to Service 5. They are not forwarded into the same shared table again, which would duplicate activity. The Service 5 HTTP ingestion endpoint is available for distinct events or a later separate-database deployment, and its actual retry behavior was verified.

## Team responsibility and Git merge ({NAMES})
Hassan Nasrallah implemented users/identity, assets and incidents. Mohammad Nasrallah implemented threats, audit logging and reports. Joseph Awaida implemented vulnerabilities, notifications and response playbooks/actions, and assembled the final integration/report/archive with Codex assistance. Original contributor names, route prefixes, code and history are preserved.
The published branches hassan/sprint-1-services-1-3 and mohammad/sprint-1-services-4-6 were fetched. Mohammad's branch contains Hassan's history. Both branches were fast-forward merged into main, retaining each original commit; no contributor history was fabricated or rewritten. The combined collection isolates runtime variables per contributor and uses Hassan's real identity endpoint for Joseph's login scenarios.

## Implementation — Services 1–3 (Hassan Nasrallah)
'''+h+'''
## Implementation — Services 4–6 (Mohammad Nasrallah)
'''+m+'''
## Implementation — Services 7–9 (Joseph Awaida)
'''+j+f'''
## Shared database design ({NAMES})
The complete SQL schema is database/schema.sql, composed in initialization order from app/schema.sql, app/sprint123_schema.sql and app/sprint456_schema.sql. The database appendix lists observed tables, column types, primary keys and foreign keys. The model includes users, assets, incidents, investigation notes, threats, audit logs, vulnerabilities, notifications, playbooks and response actions, plus compatible profile extensions, sessions, outbox/receipt tables and retained histories.
User profiles and asset profiles extend original records without deleting teammate data. Deleted-user tombstones retain identity references for audit evidence; deleted accounts cannot authenticate. Incident notes reference their incident and authenticated author. Vulnerabilities reference affected assets. Playbook actions reference their playbook and qualified assignee; ordered step numbers are unique within a playbook. Action/asset history deliberately retains original identifiers after source deletion. Database CHECK, UNIQUE and foreign-key constraints complement request validation.
Schema, seed commands, actual sanitized records, schema metadata and the runnable combined collection are included. Demo passwords are generated locally and excluded from the submission. The package supplies scripts instead of publishing a live credential-bearing database or session tokens.

## Error handling, validation and security ({NAMES})
All services use a consistent error envelope with status=error and a meaningful message; field errors include field locations and messages. Responses distinguish 401 authentication, 403 authorization/ownership, 404 absent resource, 409 uniqueness/workflow conflicts, and 422 invalid request fields. Parameterized SQL separates values from query syntax; update/search identifiers come from trusted model definitions.
Validation covers required and blank fields, email/IP/domain/hash/CVE formats, password requirements, enum values, CVSS limits, referenced identities/assets, unique indicators, ordered actions and sequential incident/vulnerability transitions. Unknown body keys are rejected. Employees cannot promote their own role; auditors cannot mutate; analysts are limited by role and assignment/ownership. Denied operations and successful mutations are audited.
Passwords use salted PBKDF2-HMAC-SHA256 with 600,000 iterations. Only token hashes are stored; sessions expire after one hour. Role/status changes revoke affected sessions, and current account status is checked on every protected request. Inactive and suspended accounts cannot log in. No raw credentials or bearer-token responses appear in public evidence.

## API documentation ({NAMES})
The API appendix lists the actual methods, routes, owners and purposes from the running application schemas. Detailed parameter, request-body, response and validation schemas are in docs/team/service-1.openapi.json through service-9.openapi.json. Each process exposes /docs. All business routes retain the implementing student's full name in their prefix.
The original three Postman collections/environments are preserved. The additional Team-All-Services collection executes all member groups in sequence against the integrated launcher, avoiding collisions between similarly named variables. A public Team-Local environment has blank password fields; prepare_team_newman writes the private populated environment only under ignored data/.

## Actual validation and evidence ({NAMES})
The merged test suite passed {validation['pytest']['tests']} tests with zero failures and errors. The actual integrated Newman run executed {validation['newman']['requests']['total']} requests with {validation['newman']['assertions']['total']} assertions and zero failures. Results are in evidence/team/pytest-results.xml, newman-output.txt, validation-summary.json and sanitized http-execution.json.
The integration verifier checked all six reports against actual database counts, required HTTP sources, verified Service 5 idempotent ingestion using real HTTP 201/200 responses, and ran the SQLite foreign-key check. This is actual execution evidence; documentation and diagrams are not presented as proof of API execution.
Genuine Postman desktop request/result screenshots and database/GitHub screenshots are included with captions identifying the service owner and purpose. Capture metadata identifies the shared team validation environment. Original Joseph evidence from 18 October is preserved as historical contribution evidence; the integrated team validation is separately dated. Token-bearing login responses use safe headers/status views, and password bodies are omitted when a secret would otherwise be visible.

## GitHub and contributions ({NAMES})
Repository: https://github.com/j123awaida/435L. The archive includes genuine commit history and original member branches are retained remotely. Contribution evidence reflects actual commits and API implementation ownership, not manufactured activity. The merge preserves the published Hassan and Mohammad commits as ancestors of main.

## References and AI disclosure ({NAMES})
Primary requirements: Final Project Fall 2026-2027-Sprint 1 (1).docx supplied for this course. Framework/reference documentation: https://fastapi.tiangolo.com/ ; https://docs.python.org/3/library/sqlite3.html ; https://docs.python.org/3/library/hashlib.html ; https://docs.pytest.org/ ; https://learning.postman.com/docs/collections/collections-overview/ .
The members' original disclosures are preserved in their contribution documents. Hassan's contribution used AI assistance for requirements/code review, implementation, tests and Postman/documentation preparation. Mohammad disclosed Claude assistance for implementation, tests, collection and documentation. Joseph used OpenAI Codex for Services 7–9, tests, execution/evidence capture and reports. Codex also assisted Joseph in the final branch merge, combined runner/collection, actual integration validation, common report and submission packaging. Every member should review and be able to explain their implementation and the disclosed assistance.

## Submission status ({NAMES})
The archive includes all nine services, database SQL/seed scripts, original and combined collections, common Word/PDF report, documentation, genuine evidence and GitHub link. Each member uploads the same common team package, as the assignment requires.
Personal and course-portal requirements remain outside software packaging: the team must confirm actual CREATES attendance and application of feedback at least three days before the true deadline; every member must review the common report and submit it through their own Moodle account. No attendance, coaching feedback, deadline or Moodle upload is invented or claimed here.
'''
    (OUT/'FINAL-TEAM-REPORT.md').write_text(text)
    return text


def screenshots():
    selected=[]
    team=ROOT/'evidence/team/screenshots/manifest.json'
    if team.exists():
        selected+=json.loads(team.read_text())
    # The original full Joseph evidence remains in the package; include a readable selection in the common report.
    old=ROOT/'evidence/screenshots/postman-manifest.json'
    if old.exists() and not team.exists():
        wanted={'Create critical vulnerability','Retrieve vulnerabilities','Retrieve vulnerability','Advance vulnerability to Resolved',
                'Search severity=Critical','Delete vulnerability','Reject analyst deletion','Create employee alert','Retrieve own notifications',
                'Retrieve specific notification','Mark own notification read','Search status=Unread','Delete notification','Reject another user notification',
                'Create playbook','List playbooks','Retrieve playbook','Add response step 1','Step 1 Completed','Delete playbook','Reject out of order action'}
        selected += [r for r in json.loads(old.read_text()) if r.get('operation') in wanted and r.get('view') in ('Response','Request')]
    old=ROOT/'evidence/screenshots/browser-manifest.json'
    if old.exists():
        wanted={'vulnerabilities-structure.png','notifications-structure.png','playbooks-structure.png','actions-structure.png',
                'vulnerability-before.png','vulnerability-after.png','notification-before.png','notification-after.png','action-before.png','action-after.png'}
        selected += [r for r in json.loads(old.read_text()) if Path(r['file']).name in wanted]
    return [r for r in selected if (ROOT/r['file']).exists()]


def api_reference():
    rows=[];schemas={}
    for number,title,owner,module,port in SERVICES:
        app=importlib.import_module('app.'+module).app;schema=app.openapi();schemas[number]=schema
        (OUT/f'service-{number}.openapi.json').write_text(json.dumps(schema,indent=2))
        for path,methods in schema['paths'].items():
            if path.endswith('/health'):continue
            for method,operation in methods.items():
                rows.append([str(number),owner,method.upper()+' '+path,operation.get('summary','')])
    return rows,schemas


def build():
    OUT.mkdir(parents=True,exist_ok=True)
    database=ROOT/'database';database.mkdir(exist_ok=True)
    (database/'schema.sql').write_text('\n\n'.join((ROOT/'app'/name).read_text() for name in ('schema.sql','sprint123_schema.sql','sprint456_schema.sql')))
    diagrams();text=narrative();api,schemas=api_reference();figures=screenshots()
    doc=Document();doc.add_heading('Security Incident Detection and Response Management System',0)
    doc.add_paragraph('Sprint I — Common Team Report\n'+NAMES+'\n20 October 2026\nhttps://github.com/j123awaida/435L')
    doc.add_heading('Table of contents',1);p=doc.add_paragraph();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'TOC \\o "1-3" \\h \\z \\u');p._p.append(field)
    doc.add_paragraph('The PDF has a populated table of contents and page numbers. Update fields in Word after any final edits.')
    styles=getSampleStyleSheet();styles.add(ParagraphStyle(name='TeamSmall',parent=styles['BodyText'],fontSize=8,leading=11,spaceAfter=5));styles['BodyText'].leading=13;styles['BodyText'].spaceAfter=7
    width=A4[0]-96
    def p(text,style='BodyText'):
        return Paragraph(escape(text),styles[style])
    story=[Spacer(1,50),p('Security Incident Detection and Response Management System','Title'),p('Sprint I — Common Team Report','Heading1'),p(NAMES),p('20 October 2026'),p('https://github.com/j123awaida/435L'),PageBreak(),p('Table of contents','Heading1')]
    toc=TableOfContents();toc.levelStyles=[ParagraphStyle(name='TeamTOC',fontSize=9,leading=13)];story += [toc,PageBreak()]
    code=False
    for line in text.splitlines():
        if line.startswith('```'):
            code=not code;continue
        if not line or line.startswith('# '):continue
        if line.startswith('## '):
            title=line[3:];doc.add_heading(title,1);story.append(p(title,'Heading1'))
        elif line.startswith('### '):
            title=line[4:];doc.add_heading(title,2);story.append(p(title,'Heading2'))
        elif line.startswith('|') and set(line.replace('|','').replace(' ','').replace(':','')) <= {'-'}:continue
        else:
            doc.add_paragraph(line);story.append(p(line,'TeamSmall' if code or line.startswith('|') else 'BodyText'))
    def picture(path,caption):
        doc.add_picture(str(path),width=Inches(6));doc.paragraphs[-1].paragraph_format.keep_with_next=True;doc.add_paragraph(caption)
        image=Image(str(path));ratio=min(width/image.imageWidth,315/image.imageHeight);image.drawWidth*=ratio;image.drawHeight*=ratio
        story.append(KeepTogether([image,p(caption,'TeamSmall'),Spacer(1,8)]))
    doc.add_heading('Architecture and database diagrams ('+NAMES+')',1);story += [PageBreak(),p('Architecture and database diagrams ('+NAMES+')','Heading1')]
    for name,label in [('architecture','Integrated service architecture'),('database-relationships','Logical database relationships')]:
        picture(OUT/(name+'.png'),NAMES+' — '+label+'. Design diagram, not execution evidence.')
    doc.add_heading('API appendix ('+NAMES+')',1);story += [PageBreak(),p('API appendix ('+NAMES+')','Heading1')]
    for number,owner,route,summary in api:
        value=f'Service {number} ({owner}): {route} — {summary}'
        doc.add_paragraph(value);story.append(p(value,'TeamSmall'))
    schema_file=ROOT/'evidence/team/database-schema.json'
    if schema_file.exists():
        doc.add_heading('Database schema appendix ('+NAMES+')',1);story += [PageBreak(),p('Database schema appendix ('+NAMES+')','Heading1')]
        for table in json.loads(schema_file.read_text()):
            value=table['name']+': '+', '.join(c['name']+' '+c['type']+(' [PK]' if c['pk'] else '') for c in table['columns'])
            doc.add_paragraph(value);story.append(p(value,'TeamSmall'))
            if table['foreign_keys']:
                value='Foreign keys: '+', '.join(f"{f['from']} -> {f['table']}.{f['to']}" for f in table['foreign_keys'])
                doc.add_paragraph(value);story.append(p(value,'TeamSmall'))
    doc.add_heading('Genuine screenshot appendix ('+NAMES+')',1);story += [PageBreak(),p('Genuine screenshot appendix ('+NAMES+')','Heading1')]
    note='Images are unedited captures from actual Postman/database/GitHub screens. Captions identify the service author and purpose. Shared team validation was operated by Joseph Awaida after the merge; historical Joseph screenshots retain their original metadata.'
    doc.add_paragraph(note);story.append(p(note))
    for figure in figures:
        picture(ROOT/figure['file'],figure['caption'])
    history=(ROOT/'evidence/team/git-history.txt')
    if history.exists():
        doc.add_heading('Actual Git history ('+NAMES+')',1);story += [PageBreak(),p('Actual Git history ('+NAMES+')','Heading1')]
        for line in history.read_text().splitlines():
            doc.add_paragraph(line);story.append(p(line,'TeamSmall'))
    doc.save(OUT/'FINAL-TEAM-REPORT.docx')
    class TeamReport(BaseDocTemplate):
        def __init__(self,filename):
            super().__init__(filename,pagesize=A4,leftMargin=48,rightMargin=48,topMargin=54,bottomMargin=48,title='Sprint I — All Nine Services',author=NAMES)
            self.addPageTemplates(PageTemplate(id='team',frames=[Frame(48,48,width,A4[1]-102,id='content')],onPage=self.footer))
        def footer(self,canvas,document):
            canvas.saveState();canvas.setFont('Helvetica',7);canvas.drawString(48,27,'Hassan Nasrallah | Mohammad Nasrallah | Joseph Awaida | Sprint I');canvas.drawRightString(A4[0]-48,27,str(document.page));canvas.restoreState()
        def afterFlowable(self,flowable):
            if isinstance(flowable,Paragraph) and flowable.style.name=='Heading1':self.notify('TOCEntry',(0,flowable.getPlainText(),self.page))
    TeamReport(str(OUT/'FINAL-TEAM-REPORT.pdf')).multiBuild(story)
    print(f'Common Word/PDF report built: all nine services, {len(api)} API operations, {len(figures)} genuine screenshots.')

if __name__=='__main__':
    build()
